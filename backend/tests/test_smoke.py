"""
End-to-end smoke test that does NOT need MySQL, Redis, or a real MRF file.

Runs against a throwaway SQLite DB + a temp DATA_ROOT, forces
EXTRACTION_MODE=stub, drives a job through the real HTTP endpoints and the
real worker code path (claim -> run_pipeline -> mark_completed), and checks
the output zip + KPI aggregates.

    cd backend
    python -m pytest -q
"""
from __future__ import annotations

import gzip
import importlib
import io
import sys
import zipfile
from pathlib import Path

import pytest


@pytest.fixture()
def app_ctx(tmp_path, monkeypatch):
    db_file = tmp_path / "test.db"
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_file.as_posix()}")
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "data"))
    monkeypatch.setenv("EXTRACTION_MODE", "stub")
    monkeypatch.setenv("CORS_ORIGINS", "http://localhost:5173")

    # drop every app.* module so they re-import fresh against this env
    # (a clean Base/MetaData each time -> no "table already defined")
    for name in [n for n in list(sys.modules) if n == "app" or n.startswith("app.")]:
        del sys.modules[name]

    database = importlib.import_module("app.database")
    database.init_db()
    main = importlib.import_module("app.main")

    from fastapi.testclient import TestClient

    with TestClient(main.app) as client:
        yield client, database, tmp_path


def _fake_mrf_gz() -> bytes:
    buf = io.BytesIO()
    with gzip.GzipFile(fileobj=buf, mode="wb") as gz:
        gz.write(b'{"in_network": [], "provider_references": []}')
    return buf.getvalue()


def test_health(app_ctx):
    client, *_ = app_ctx
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json()["extraction_mode"] == "stub"


def test_companies(app_ctx):
    client, *_ = app_ctx
    r = client.get("/api/companies")
    assert r.status_code == 200
    keys = {c["key"] for c in r.json()}
    assert {"uhc", "aetna", "mvp", "anthem", "cigna"} <= keys


def test_full_pipeline_stub(app_ctx):
    client, database, tmp_path = app_ctx

    # 1. chunked upload
    payload = _fake_mrf_gz()
    init = client.post(
        "/api/jobs/uploads",
        json={"insurance_company": "aetna", "filename": "aetna_rates.json.gz",
              "file_size": len(payload)},
    )
    assert init.status_code == 200, init.text
    upload_id = init.json()["upload_id"]

    mid = len(payload) // 2
    for idx, chunk in enumerate((payload[:mid], payload[mid:])):
        r = client.post(
            f"/api/jobs/uploads/{upload_id}/chunks",
            data={"index": idx},
            files={"chunk": ("blob", chunk, "application/octet-stream")},
        )
        assert r.status_code == 200, r.text

    fin = client.post(f"/api/jobs/uploads/{upload_id}/complete", json={"total_chunks": 2})
    assert fin.status_code == 200, fin.text
    job = fin.json()
    job_id = job["id"]
    assert job["status"] == "queued"
    assert job["total_bytes"] == len(payload)

    # raw file assembled on disk
    assert Path(job["source_filename"]).suffix == ".gz"

    # 2. run the worker's single-job path (stub pipeline, ~a few seconds)
    import app.worker as worker

    assert worker.process_one() is True

    # 3. job completed with output
    detail = client.get(f"/api/jobs/{job_id}").json()
    assert detail["status"] == "completed", detail
    assert detail["record_count"] > 0
    assert detail["output_file_count"] >= 1
    assert detail["has_output"] is True
    assert detail["percent_complete"] == 100.0

    # 4. status endpoint agrees
    st = client.get(f"/api/jobs/{job_id}/status").json()
    assert st["status"] == "completed"
    assert st["percent"] == 100.0

    # 5. download: the default is a single Excel workbook, named descriptively
    xl = client.get(f"/api/jobs/{job_id}/download")
    assert xl.status_code == 200, xl.text
    assert xl.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="AETNA_' in xl.headers["content-disposition"]
    assert xl.headers["content-disposition"].rstrip('"').endswith(".xlsx")
    assert detail["has_xlsx"] is True
    assert detail["download_name"].startswith("AETNA_") and detail["download_name"].endswith(".xlsx")

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xl.content), read_only=True)
    assert "rates" in wb.sheetnames and "summary" in wb.sheetnames
    rows = list(wb["rates"].iter_rows(values_only=True))
    assert rows[0][0] == "billing_code"
    assert len(rows) - 1 == detail["record_count"]

    # ...and the CSV zip is still there on request
    dl = client.get(f"/api/jobs/{job_id}/download?format=zip")
    assert dl.status_code == 200
    assert dl.headers["content-type"] == "application/zip"
    assert dl.headers["content-disposition"].rstrip('"').endswith(".zip")
    zf = zipfile.ZipFile(io.BytesIO(dl.content))
    names = zf.namelist()
    assert any(n.startswith("part_") and n.endswith(".csv") for n in names)
    first_csv = next(n for n in names if n.endswith(".csv"))
    header = zf.read(first_csv).splitlines()[0].decode()
    assert "billing_code" in header
    # zip columns the data team needs must be present
    for col in ("provider_zips", "provider_zip5s", "primary_zip"):
        assert col in header, header

    # 6. list + filters
    lst = client.get("/api/jobs", params={"company": "aetna", "status": "completed"}).json()
    assert lst["total"] >= 1
    assert all(j["insurance_company"] == "aetna" for j in lst["jobs"])

    # 7. KPIs reflect the completed job
    kpi = client.get("/api/kpis").json()
    assert kpi["totals"]["total_jobs"] >= 1
    assert kpi["totals"]["completed_jobs"] >= 1
    assert kpi["totals"]["total_records"] >= detail["record_count"]
    assert kpi["totals"]["success_rate"] == pytest.approx(1.0)
    assert any(c["company"] == "aetna" for c in kpi["by_company"])
    assert len(kpi["volume_over_time"]) == 31


def test_bad_company_rejected(app_ctx):
    client, *_ = app_ctx
    r = client.post(
        "/api/jobs/uploads",
        json={"insurance_company": "notapayer", "filename": "x.json.gz",
              "file_size": 10},
    )
    assert r.status_code == 422


def test_bad_extension_rejected(app_ctx):
    client, *_ = app_ctx
    r = client.post(
        "/api/jobs/uploads",
        json={"insurance_company": "uhc", "filename": "notmrf.txt", "file_size": 10},
    )
    assert r.status_code == 422
