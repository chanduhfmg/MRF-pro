"""Jobs finished before the workbook existed get one built on first download."""
from __future__ import annotations

import io

from tests.test_search import TARGET_NPI, _run_search, app_ctx  # noqa: F401  (fixture)


def test_xlsx_is_built_lazily_from_csv_parts(app_ctx):
    client, _ = app_ctx
    s = _run_search(client, search_type="npi", value=TARGET_NPI)
    job = client.post(f"/api/searches/{s['id']}/extract", json={"top_n": 1}).json()["jobs"][0]

    import app.worker as worker
    assert worker.process_one() is True

    # simulate a pre-feature job: no workbook on disk, no path on the row
    from pathlib import Path

    from app import crud
    from app.database import session_scope

    with session_scope() as db:
        row = crud.get_job(db, job["id"])
        Path(row.output_xlsx_path).unlink()
        row.output_xlsx_path = None
    assert client.get(f"/api/jobs/{job['id']}").json()["has_xlsx"] is False

    r = client.get(f"/api/jobs/{job['id']}/download")  # default = xlsx
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    rows = list(wb["rates"].iter_rows(values_only=True))
    assert len(rows) - 1 == client.get(f"/api/jobs/{job['id']}").json()["record_count"]

    # and it's remembered, not rebuilt every time
    assert client.get(f"/api/jobs/{job['id']}").json()["has_xlsx"] is True
