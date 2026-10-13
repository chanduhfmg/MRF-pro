"""
End-to-end run of the REAL pipeline on a synthetic MRF, with NPPES mocked.

real.py carries the phase plan, the uncapped enrichment, the provider
directory and the workbook - all of which the stub-mode tests skip entirely.
"""
from __future__ import annotations

import csv
import gzip
import importlib
import io
import json
import sys
import zipfile

import pytest

TIN = "473943680"
# 9 NPIs in one group: more than the old 5-per-row cap, so a capped run
# would leave four of them with no ZIP at all.
NPIS = [str(1043703325 + i) for i in range(9)]
OTHER_TIN = "111111111"
OTHER_NPIS = [str(1900000000 + i) for i in range(3)]


def _mrf_bytes() -> bytes:
    doc = {
        "reporting_entity_name": "Test Payer",
        "provider_references": [
            {"provider_group_id": 1,
             "provider_groups": [{"npi": NPIS, "tin": {"type": "ein", "value": TIN}}]},
            {"provider_group_id": 2,
             "provider_groups": [{"npi": OTHER_NPIS, "tin": {"type": "ein", "value": OTHER_TIN}}]},
        ],
        "in_network": [
            {
                "billing_code": code, "billing_code_type": "CPT", "name": f"Service {code}",
                "negotiated_rates": [
                    {"provider_references": [1, 2],
                     "negotiated_prices": [{"negotiated_type": "negotiated",
                                            "negotiated_rate": 100 + i,
                                            "billing_class": "professional"}]}
                ],
            }
            for i, code in enumerate(["99213", "99214", "92014"])
        ],
    }
    buf = io.BytesIO()
    with gzip.GzipFile(fileobj=buf, mode="wb") as gz:
        gz.write(json.dumps(doc).encode())
    return buf.getvalue()


@pytest.fixture()
def real_ctx(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{(tmp_path / 'real.db').as_posix()}")
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "data"))
    monkeypatch.setenv("EXTRACTION_MODE", "real")
    monkeypatch.setenv("NPI_MAX_PER_ROW", "0")     # every NPI
    monkeypatch.setenv("NPI_LOOKUP_THREADS", "2")

    for name in [n for n in list(sys.modules) if n == "app" or n.startswith("app.")]:
        del sys.modules[name]
    database = importlib.import_module("app.database")
    database.init_db()
    main = importlib.import_module("app.main")

    from app.config import REPO_ROOT
    if str(REPO_ROOT) not in sys.path:
        sys.path.insert(0, str(REPO_ROOT))
    import enrich_npi_all as en

    looked_up: list[str] = []

    def fake_fetch(npi):
        looked_up.append(npi)
        tail = int(npi) % 100
        return {
            "provider_name": f"Dr {npi[-4:]}", "organization_name": "",
            "address": f"{tail} Main St", "city": "Rochester", "state": "NY",
            # a ZIP+4 whose 5-digit form differs per provider: exactly what an
            # analyst filters on
            "zip": f"14{tail:03d}1234"[:9], "zip5": f"14{tail:03d}",
        }

    monkeypatch.setattr(en, "_fetch_npi", fake_fetch)

    from fastapi.testclient import TestClient
    with TestClient(main.app) as client:
        yield client, tmp_path, looked_up


def _run(client, tmp_path, **job_kw):
    from app import crud
    from app.database import session_scope

    src = tmp_path / "in.json.gz"
    src.write_bytes(_mrf_bytes())
    with session_scope() as db:
        job = crud.create_job(
            db, insurance_company="cigna", source_filename=src.name,
            source_file_path=str(src), total_bytes=src.stat().st_size, **job_kw,
        )
        jid = job.id

    import app.worker as worker
    assert worker.process_one() is True
    d = client.get(f"/api/jobs/{jid}").json()
    assert d["status"] == "completed", d.get("error_message")
    return jid, d


def test_real_run_enriches_every_npi_and_writes_a_provider_directory(real_ctx):
    client, tmp_path, looked_up = real_ctx
    jid, d = _run(
        client, tmp_path,
        target_type="tin", target_value=TIN, filter_to_target=True,
        plan_name="Cigna Health Life Insurance Company PathWell PPO",
        code_set_name="Cardiology",
        target_codes=json.dumps(["99213", "99214", "92014"]),
    )

    # 3 codes x 1 matching group = 3 rows, all for the searched TIN
    assert d["record_count"] == 3

    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/jobs/{jid}/download?format=zip").content))
    part_name = next(n for n in z.namelist() if n.endswith("part_0001.csv"))
    assert part_name == "CIGNA_Cardiology_part_0001.csv"
    rates = list(csv.DictReader(io.StringIO(z.read(part_name).decode())))
    assert {r["tin"] for r in rates} == {TIN}

    # every one of the 9 NPIs has its own ZIP on the row - not just the first 5
    for r in rates:
        assert r["npis"].split(";") == NPIS
        zip5s = r["provider_zip5s"].split(";")
        assert len(zip5s) == 9 and all(zip5s), zip5s
        assert len(set(zip5s)) == 9, "each provider keeps its own ZIP"
        assert r["primary_zip"] == zip5s[0]

    # ...and one line per NPI for filtering
    providers = list(csv.DictReader(io.StringIO(z.read("providers.csv").decode())))
    assert {p["npi"] for p in providers} == set(NPIS)
    assert all(p["zip5"] and p["address"] and p["city"] for p in providers)
    assert all(p["tins"] == TIN for p in providers)
    assert all(int(p["row_count"]) == 3 for p in providers)
    # off-target providers were filtered out, so they are not in the directory
    assert not ({p["npi"] for p in providers} & set(OTHER_NPIS))

    # each unique NPI costs exactly one NPPES call however many rows use it
    assert sorted(looked_up) == sorted(NPIS)

    # workbook carries both sheets, and the filename says what it is
    from openpyxl import load_workbook
    r = client.get(f"/api/jobs/{jid}/download")
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert {"rates", "providers", "summary"} <= set(wb.sheetnames)
    assert len(list(wb["providers"].iter_rows(values_only=True))) - 1 == 9
    assert "CIGNA_Cigna-Life-PathWell_TIN473943680_" in r.headers["content-disposition"]
    assert "_Cardiology.xlsx" in r.headers["content-disposition"]


def test_real_run_unfiltered_keeps_every_provider_group(real_ctx):
    client, tmp_path, looked_up = real_ctx
    jid, d = _run(client, tmp_path, target_type="tin", target_value=TIN, filter_to_target=False)

    assert d["record_count"] == 6          # 3 codes x 2 groups
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/jobs/{jid}/download?format=zip").content))
    providers = list(csv.DictReader(io.StringIO(z.read("providers.csv").decode())))
    assert {p["npi"] for p in providers} == set(NPIS) | set(OTHER_NPIS)
    assert sorted(looked_up) == sorted(NPIS + OTHER_NPIS)
    # ALL, not the id, because the run was not narrowed
    assert "_ALL_" in client.get(f"/api/jobs/{jid}/download").headers["content-disposition"]


def test_real_run_manifest_and_phase_timings(real_ctx):
    client, tmp_path, _ = real_ctx
    jid, _ = _run(client, tmp_path, target_type="tin", target_value=TIN, filter_to_target=True)

    m = client.get(f"/api/jobs/{jid}/manifest").json()
    assert m["row_count"] == 3
    assert m["provider_count"] == 9
    assert m["providers_csv"] == "providers.csv"
    assert m["npi_max_per_row"] == 0
    assert sorted(m["extract_stats"]["matched_codes"]) == ["92014", "99213", "99214"]
    assert m["extract_stats"]["dropped_off_target"] == 3      # the other group, 3 codes
    # every phase finished before packaging recorded a duration (including a
    # phase that took ~0s - absent would mean "never ran")
    assert set(m["phase_seconds"]) == {"lookup", "extract"}
    assert all(v >= 0 for v in m["phase_seconds"].values())


def test_real_run_progress_is_monotonic_and_eta_is_sane(real_ctx, monkeypatch):
    """The old model let percent jump backwards (94% -> 56%) and reported
    ~1s remaining for an hours-long job. Record every progress write and
    assert neither can happen."""
    client, tmp_path, _ = real_ctx

    from app import crud

    seen: list[tuple[float, float | None]] = []
    real_update = crud.update_progress

    def spy(db, job_id, **kw):
        real_update(db, job_id, **kw)
        job = crud.get_job(db, job_id)
        seen.append((job.percent_complete, job.eta_seconds_est))

    monkeypatch.setattr(crud, "update_progress", spy)
    _run(client, tmp_path, target_type="tin", target_value=TIN, filter_to_target=True)

    assert seen, "the pipeline must report progress"
    percents = [p for p, _ in seen]
    for a, b in zip(percents, percents[1:]):
        assert b >= a - 0.11, f"percent went backwards: {a} -> {b}"
    assert percents[-1] >= 99.0
    etas = [e for _, e in seen if e is not None]
    assert etas, "an ETA must be published"
    assert all(e >= 0 for e in etas)
