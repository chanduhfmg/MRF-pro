"""
The time-based phase plan (accurate, monotonic, countdown-friendly ETA) and
the provider directory (one line per NPI, every NPI, in CSV zip + workbook).
"""
from __future__ import annotations

import csv
import io
import time
import zipfile

from tests.test_search import TARGET_NPI, _run_search, app_ctx  # noqa: F401  (fixture)


# --- PhasePlan --------------------------------------------------------------

def test_phase_plan_percent_is_time_proportional_and_never_goes_backwards():
    from app.eta import PhasePlan

    plan = PhasePlan([("download", 10), ("lookup", 30), ("extract", 60)])
    assert plan.percent() == 0.0

    plan.start("download")
    plan.observe(0.5)
    assert plan.percent() == 5.0            # 5s of a 100s plan
    plan.observe(0.2)                        # a stale/lower sample must not regress
    assert plan.percent() == 5.0
    plan.finish("download")                  # actual duration (~0s) replaces the estimate
    # download took ~0s, so the plan is now ~90s and nothing of it is done
    assert plan.percent() < 2.0

    plan.start("lookup")
    plan.observe(1.0)
    p_lookup_done = plan.percent()
    plan.finish("lookup")
    plan.start("extract")
    plan.observe(0.5)
    assert plan.percent() > p_lookup_done   # monotonic across phase boundaries
    plan.observe(1.0)
    plan.finish("extract")
    assert plan.percent() == 100.0
    assert plan.eta_seconds() == 0.0


def test_phase_plan_eta_counts_remaining_phases_and_uses_live_rate():
    from app.eta import PhasePlan

    plan = PhasePlan([("lookup", 20), ("extract", 40)])
    plan.start("lookup")
    # before any progress: prior estimate for this phase + the next phase
    assert 59 <= plan.eta_seconds() <= 60
    plan.observe(0.5)
    eta = plan.eta_seconds()
    # half the lookup done (prior says ~10s left) + full extract (40s)
    assert 45 <= eta <= 55, eta
    plan.finish("lookup")
    plan.start("extract")
    assert 39 <= plan.eta_seconds() <= 41


def test_phase_plan_live_rate_dominates_once_a_phase_has_run(monkeypatch):
    from app import eta as eta_mod
    from app.eta import PhasePlan

    # one fake clock for every timing read in the module (see eta._now)
    now = {"t": 1000.0}
    monkeypatch.setattr(eta_mod, "_now", lambda: now["t"])

    plan = PhasePlan([("extract", 1000)])   # prior says 1000s
    plan.start("extract")
    # reality: 50% done after 60s -> live rate says ~60s left, not ~500s
    now["t"] += 60
    plan.observe(0.5)
    eta = plan.eta_seconds()
    assert eta < 150, f"live rate should dominate after 60s, got {eta}"

    # and a phase that is genuinely slower than its prior is reported as such
    now["t"] += 600          # 660s elapsed, still only 60% done
    plan.observe(0.6)
    assert plan.eta_seconds() > 300


def test_learned_phase_rates_persist(app_ctx):
    from app import eta as eta_mod

    before = eta_mod.load_phase_rates()["lookup"]
    eta_mod.learn_phase_rate("lookup", size_bytes=100 * 1024 * 1024, seconds=100.0)  # 1 MB/s
    after = eta_mod.load_phase_rates()["lookup"]
    assert after != before
    assert after < before, "a slower observed run pulls the estimate down"
    # ignores nonsense
    eta_mod.learn_phase_rate("lookup", size_bytes=0, seconds=10.0)
    assert eta_mod.load_phase_rates()["lookup"] == after


# --- provider directory ------------------------------------------------------

def test_output_includes_one_line_per_npi_in_zip_and_workbook(app_ctx):
    client, _ = app_ctx
    s = _run_search(client, search_type="npi", value=TARGET_NPI)
    job = client.post(f"/api/searches/{s['id']}/extract", json={"top_n": 1}).json()["jobs"][0]

    import app.worker as worker
    assert worker.process_one() is True
    detail = client.get(f"/api/jobs/{job['id']}").json()
    assert detail["status"] == "completed"

    # zip: rates parts + providers.csv + manifest
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/jobs/{job['id']}/download?format=zip").content))
    names = z.namelist()
    assert "providers.csv" in names and "manifest.json" in names
    providers = list(csv.DictReader(io.StringIO(z.read("providers.csv").decode())))
    assert providers, "directory must not be empty"
    assert {"npi", "zip", "zip5", "tins", "row_count", "address"} <= set(providers[0].keys())
    # every NPI on every rate row is in the directory
    rate_npis: set[str] = set()
    for n in names:
        if n.startswith("part_"):
            for row in csv.DictReader(io.StringIO(z.read(n).decode())):
                rate_npis.update(x for x in row["npis"].split(";") if x)
    assert rate_npis == {p["npi"] for p in providers}
    assert all(p["zip5"] and len(p["zip5"]) == 5 for p in providers)

    # workbook: a 'providers' sheet with the same rows
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job['id']}/download").content), read_only=True)
    assert "providers" in wb.sheetnames
    rows = list(wb["providers"].iter_rows(values_only=True))
    assert rows[0][0] == "npi"
    assert len(rows) - 1 == len(providers)
    # ZIPs stay text (leading zeros survive), row_count is a number
    hdr = list(rows[0])
    assert isinstance(rows[1][hdr.index("zip5")], str)
    assert isinstance(rows[1][hdr.index("row_count")], int)


def test_uncapped_enrichment_is_the_shipped_default():
    """0 = every NPI. Asserted on the Settings class default rather than the
    live .env, so a stale local override can't make this pass or fail."""
    from app.config import Settings

    assert Settings.model_fields["npi_max_per_row"].default == 0


def test_enrichment_is_uncapped_by_default_and_fills_every_npi(monkeypatch):
    from app.config import REPO_ROOT
    import sys

    if str(REPO_ROOT) not in sys.path:
        sys.path.insert(0, str(REPO_ROOT))
    import enrich_npi_all as en

    # 12 NPIs on one row; the old cap of 5 would leave 7 without a ZIP
    npis = [str(1000000000 + i) for i in range(12)]
    row = {"tin": "123456789", "npis": ";".join(npis)}
    calls: list[str] = []

    def fake_fetch(npi):
        calls.append(npi)
        return {**en.EMPTY_NPI_INFO, "zip": f"1460{int(npi) % 10}3923", "zip5": f"1460{int(npi) % 10}"}

    monkeypatch.setattr(en, "_fetch_npi", fake_fetch)
    cache: dict = {}
    out = list(en.iter_enriched_rows(iter([row]), threads=1, max_npis_per_row=0, cache=cache))
    assert len(out) == 1
    zips = out[0]["provider_zip5s"].split(";")
    assert len(zips) == 12 and all(zips), zips
    assert sorted(calls) == sorted(npis)          # every NPI looked up exactly once
    assert set(cache) == set(npis)                # and the caller gets the directory

    # the explicit cap still works when someone asks for it
    calls.clear()
    out = list(en.iter_enriched_rows(iter([dict(row)]), threads=1, max_npis_per_row=5))
    assert len(out[0]["provider_zip5s"].split(";")) == 5


def test_xlsx_cells_over_excel_limit_are_truncated_not_rejected(tmp_path):
    from app.pipeline.package import XLSX_MAX_CELL, write_split_csvs, write_xlsx

    huge = ";".join(str(1000000000 + i) for i in range(4000))  # ~44k chars
    parts, n = write_split_csvs(iter([{"billing_code": "99213", "npis": huge, "npi_count": "4000"}]),
                                tmp_path, ["billing_code", "npis", "npi_count"])
    path = write_xlsx(parts, tmp_path, ["billing_code", "npis", "npi_count"], n)
    from openpyxl import load_workbook
    rows = list(load_workbook(path, read_only=True)["rates"].iter_rows(values_only=True))
    cell = rows[1][1]
    assert len(cell) <= XLSX_MAX_CELL
    assert cell.endswith("providers sheet / CSV)")
    assert rows[1][2] == 4000


def test_rate_estimator_decay_is_time_weighted_not_sample_weighted(monkeypatch):
    """A long slow stretch must outweigh a short fast one.

    Sample-counted decay reported a cheerful ETA for a job that had been
    crawling for ten minutes, because the one fast sample from the start
    still carried equal weight.
    """
    from app import eta as eta_mod
    from app.eta import RateEstimator

    now = {"t": 0.0}
    monkeypatch.setattr(eta_mod, "_now", lambda: now["t"])

    est = RateEstimator()
    now["t"] += 60
    est.observe(0.5)                       # fast: half in the first minute
    fast_rate = est.rate

    now["t"] += 600
    est.observe(0.6)                       # then a crawl: 10% in ten minutes
    assert est.rate < fast_rate / 5, "ten slow minutes must dominate one fast one"

    # remaining 40% at the (now slow) rate is a long wait, and honest about it
    assert est.eta_seconds(1.0) > 300
