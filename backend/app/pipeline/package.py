"""
Shared output packaging: take an iterable of dict rows, write them into
CSV parts that roll over at CSV_ROW_THRESHOLD rows, then zip every part
into {job_id}.zip. Used by both the stub and the real pipeline.
"""
from __future__ import annotations

import csv
import logging
import sys
import zipfile
from collections.abc import Iterable, Iterator
from pathlib import Path

from ..config import settings

log = logging.getLogger("mrf.worker.package")

# some MRF rows have huge npis fields; mirror enrich_npi_all.py's bump
try:
    csv.field_size_limit(sys.maxsize)
except OverflowError:
    csv.field_size_limit(2**31 - 1)


class CsvSplitWriter:
    """Writes dict rows to {prefix}part_0001.csv, {prefix}part_0002.csv, ...
    rolling over every `row_threshold` rows. Each part carries the full
    header. Flushes after every row so an interrupted run keeps everything
    already written (same guarantee the original scripts give).

    `max_parts` is a hard ceiling: once reached, the writer stops opening
    new files and keeps appending to the last one instead of rolling over,
    so a job never produces more than `max_parts` CSVs no matter how many
    rows it has."""

    def __init__(
        self,
        out_dir: Path,
        fieldnames: list[str],
        row_threshold: int,
        *,
        max_parts: int | None = None,
        prefix: str = "",
    ):
        self.out_dir = out_dir
        self.fieldnames = fieldnames
        self.row_threshold = max(1, row_threshold)
        self.max_parts = max(1, max_parts) if max_parts else None
        self.prefix = prefix
        self.part_index = 0
        self.rows_in_part = 0
        self.total_rows = 0
        self._fh = None
        self._writer = None
        self.parts: list[Path] = []

    def _open_next(self) -> None:
        self._close()
        self.part_index += 1
        path = self.out_dir / f"{self.prefix}part_{self.part_index:04d}.csv"
        self._fh = open(path, "w", newline="", encoding="utf-8")
        self._writer = csv.DictWriter(self._fh, fieldnames=self.fieldnames, extrasaction="ignore")
        self._writer.writeheader()
        self._fh.flush()
        self.parts.append(path)
        self.rows_in_part = 0

    def write(self, row: dict) -> None:
        at_cap = self.max_parts is not None and self.part_index >= self.max_parts
        if self._writer is None or (self.rows_in_part >= self.row_threshold and not at_cap):
            self._open_next()
        self._writer.writerow({k: row.get(k, "") for k in self.fieldnames})
        self._fh.flush()
        self.rows_in_part += 1
        self.total_rows += 1

    def _close(self) -> None:
        if self._fh is not None:
            self._fh.close()
            self._fh = None
            self._writer = None

    def close(self) -> None:
        self._close()

    def __enter__(self) -> "CsvSplitWriter":
        return self

    def __exit__(self, *exc) -> None:
        self._close()


def write_split_csvs(
    rows: Iterable[dict],
    out_dir: Path,
    fieldnames: list[str],
    *,
    row_threshold: int | None = None,
    max_parts: int | None = None,
    prefix: str = "",
    on_row: "callable | None" = None,
) -> tuple[list[Path], int]:
    """Returns (csv_part_paths, total_rows)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    threshold = row_threshold or settings.csv_row_threshold
    cap = settings.csv_max_parts if max_parts is None else max_parts
    with CsvSplitWriter(out_dir, fieldnames, threshold, max_parts=cap, prefix=prefix) as w:
        for row in rows:
            w.write(row)
            if on_row is not None:
                on_row(w.total_rows)
    return w.parts, w.total_rows


def zip_parts(job_id: str, out_dir: Path, part_paths: list[Path]) -> tuple[Path, int]:
    """Zip the CSV parts (plus manifest.json and providers.csv when present)
    into {job_id}.zip. Returns (zip_path, zip_size_bytes)."""
    zip_path = out_dir / f"{job_id}.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for p in part_paths:
            if p.exists():
                zf.write(p, arcname=p.name)
        for extra in ("manifest.json", PROVIDERS_CSV):
            f = out_dir / extra
            if f.exists():
                zf.write(f, arcname=extra)
    return zip_path, zip_path.stat().st_size


# --- provider directory ------------------------------------------------------
#
# The rate rows carry every NPI's details as ';'-joined lists aligned with the
# `npis` column. That is complete but hostile to "which provider is at ZIP
# 07202?" - so the same information is also written one line per NPI.

PROVIDERS_CSV = "providers.csv"
PROVIDER_COLUMNS = [
    "npi", "provider_name", "organization_name", "address", "city", "state",
    "zip", "zip5", "tins", "row_count",
]


def write_providers_csv(out_dir: Path, directory: dict, npi_tins: dict, npi_rows: dict) -> Path | None:
    """directory: npi -> NPPES info (from the enrich cache); npi_tins: npi ->
    set of TINs it appeared under; npi_rows: npi -> number of rate rows."""
    if not directory:
        return None
    path = out_dir / PROVIDERS_CSV
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=PROVIDER_COLUMNS)
        w.writeheader()
        for npi in sorted(directory, key=lambda n: (-(npi_rows.get(n, 0)), n)):
            info = directory[npi] or {}
            w.writerow({
                "npi": npi,
                "provider_name": info.get("provider_name", ""),
                "organization_name": info.get("organization_name", ""),
                "address": info.get("address", ""),
                "city": info.get("city", ""),
                "state": info.get("state", ""),
                "zip": info.get("zip", ""),
                "zip5": info.get("zip5", ""),
                "tins": ";".join(sorted(npi_tins.get(npi, ()))),
                "row_count": npi_rows.get(npi, 0),
            })
    return path


def track_npis(rows, npi_tins: dict, npi_rows: dict):
    """Pass-through generator that records, for every enriched row, which
    TIN each NPI sat under and how many rows it appears on."""
    for row in rows:
        tin = row.get("tin") or ""
        for n in (row.get("npis") or "").split(";"):
            n = n.strip()
            if not n:
                continue
            npi_rows[n] = npi_rows.get(n, 0) + 1
            if tin:
                npi_tins.setdefault(n, set()).add(str(tin))
        yield row


def iter_csv_rows(path: Path) -> Iterator[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        yield from csv.DictReader(f)


# Excel's hard limit is 1,048,576 rows per sheet, header included. Stay
# comfortably under it; above this the CSV zip is the only deliverable.
XLSX_MAX_ROWS = 1_000_000

# Columns Excel should treat as numbers. Everything else stays text on
# purpose - NPIs, TINs and ZIPs have leading zeros that a numeric cell would
# silently destroy.
_NUMERIC_COLUMNS = {"negotiated_rate", "npi_count"}


# Excel refuses cells longer than this; an uncapped provider group can have
# thousands of NPIs and blow past it in the ';'-joined columns.
XLSX_MAX_CELL = 32_767
_TRUNCATED = " …(truncated: full list in providers sheet / CSV)"


def write_xlsx(
    part_paths: list[Path],
    out_dir: Path,
    fieldnames: list[str],
    total_rows: int,
    *,
    summary: dict | None = None,
    providers_csv: Path | None = None,
    on_row=None,
) -> Path | None:
    """One workbook: 'rates' (every extracted row), 'providers' (one line per
    NPI, when providers.csv exists) and 'summary'. Returns None when the
    output is too large for a sheet, so the caller can fall back to the zip.

    Uses openpyxl's write-only mode: rows stream to disk instead of being
    held in memory, which matters at hundreds of thousands of rows.
    """
    if total_rows > XLSX_MAX_ROWS:
        log.info("skipping xlsx: %d rows exceeds the %d-row sheet limit", total_rows, XLSX_MAX_ROWS)
        return None

    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("rates")
    ws.append(fieldnames)
    n = 0
    for p in part_paths:
        if not p.exists():
            continue
        for row in iter_csv_rows(p):
            ws.append([_cell(k, row.get(k, "")) for k in fieldnames])
            n += 1
            if on_row is not None and n % 5000 == 0:
                on_row(n)

    if providers_csv is not None and providers_csv.exists():
        ps = wb.create_sheet("providers")
        ps.append(PROVIDER_COLUMNS)
        for row in iter_csv_rows(providers_csv):
            ps.append([_cell(k, row.get(k, "")) for k in PROVIDER_COLUMNS])

    if summary:
        s = wb.create_sheet("summary")
        s.append(["field", "value"])
        for k, v in summary.items():
            s.append([k, v if isinstance(v, (int, float, str)) or v is None else str(v)])

    path = out_dir / "output.xlsx"
    wb.save(path)
    return path


def _cell(column: str, value: str):
    if column in _NUMERIC_COLUMNS or column == "row_count":
        if value not in ("", None):
            try:
                f = float(value)
                return int(f) if f.is_integer() and column != "negotiated_rate" else f
            except ValueError:
                return value
        return value
    if isinstance(value, str) and len(value) > XLSX_MAX_CELL:
        return value[: XLSX_MAX_CELL - len(_TRUNCATED)] + _TRUNCATED
    return value
